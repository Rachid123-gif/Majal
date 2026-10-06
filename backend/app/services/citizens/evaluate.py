"""Evaluation of the analysis (BRIEF §9.6): precision and recall of the theme classification,
accuracy of tonality, language and location — for the local model and for the keyword fallback.

Two bases, always stated with the figures:
- provisional: the annotation written by Claude, who also wrote the fictitious texts (partly
  circular);
- reference: the contributions classified by the professor (docs/evaluation/
  annotation-professeur.xlsx), Amazigh texts excluded (approximate transcriptions).
"""

from collections.abc import Iterable
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook

from app.models import Contribution


def _prf(pairs: Iterable[tuple[set[str], set[str]]]) -> dict[str, Any]:
    tp = fp = fn = 0
    for predicted, truth in pairs:
        tp += len(predicted & truth)
        fp += len(predicted - truth)
        fn += len(truth - predicted)
    precision = tp / (tp + fp) if tp + fp else None
    recall = tp / (tp + fn) if tp + fn else None
    f1 = (
        2 * precision * recall / (precision + recall)
        if precision is not None and recall is not None and precision + recall
        else None
    )
    return {"precision": precision, "recall": recall, "f1": f1, "tp": tp, "fp": fp, "fn": fn}


def _accuracy(pairs: list[tuple[Any, Any]]) -> dict[str, Any]:
    scored = [(p, t) for p, t in pairs if t not in (None, "")]
    correct = sum(1 for p, t in scored if p == t)
    return {
        "accuracy": correct / len(scored) if scored else None,
        "n": len(scored),
        "correct": correct,
    }


def score(contributions: list[Contribution], truth: dict[str, dict[str, Any]]) -> dict[str, Any]:
    rows = [(c, truth[c.external_id]) for c in contributions if c.external_id in truth]
    ai_rows = [(c, t) for c, t in rows if c.analysis.get("mode") == "ai"]
    result: dict[str, Any] = {
        "n": len(rows),
        "n_ai": len(ai_rows),
        "themes": {
            "model": _prf((set(c.themes), set(t["themes"])) for c, t in rows),
            "keywords": _prf(
                (set(c.analysis.get("keywords", {}).get("themes", [])), set(t["themes"]))
                for c, t in rows
            ),
        },
        # Main theme (the one the statistics count): is it among the annotated themes?
        "main_theme": {
            "model": _accuracy(
                [
                    (
                        t["themes"][0] if c.themes and c.themes[0] in t["themes"] else None,
                        t["themes"][0],
                    )
                    for c, t in rows
                ]
            )
            if all(t["themes"] for _, t in rows)
            else None,
        },
        "tonality": {
            "model": _accuracy([(c.tonality, t.get("tonality")) for c, t in rows]),
            "keywords": _accuracy(
                [
                    (c.analysis.get("keywords", {}).get("tonality"), t.get("tonality"))
                    for c, t in rows
                ]
            ),
        },
    }
    if all("language" in t for _, t in rows):
        result["language"] = {
            "model": _accuracy([(c.language, t["language"]) for c, t in rows]),
            "keywords": _accuracy(
                [(c.analysis.get("keywords", {}).get("language"), t["language"]) for c, t in rows]
            ),
        }
    if all("territory_id" in t for _, t in rows):
        located = [(c, t) for c, t in rows if c.territory_id is not None]
        result["location"] = {
            "located": len(located),
            "correct": sum(1 for c, t in located if c.territory_id == t["territory_id"]),
            "wrong": sum(1 for c, t in located if c.territory_id != t["territory_id"]),
            "not_located": len(rows) - len(located),
        }
    return result


def provisional_truth(path: Path, unit_ids: dict[str, int]) -> dict[str, dict[str, Any]]:
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    return {
        c["id"]: {
            "themes": c["themes"],
            "tonality": c["tonality"],
            "language": c["language"],
            "territory_id": unit_ids.get(c["unit"]),
        }
        for c in data["contributions"]
    }


def reference_truth(path: Path, excluded_languages: set[str]) -> dict[str, dict[str, Any]]:
    """The professor's classification. Rows left empty (or in an excluded language) are ignored."""
    if not path.exists():
        return {}
    sheet = load_workbook(path, read_only=True, data_only=True)["À classer"]
    truth = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        _, identifier, language, _, themes, tonality, _place = (list(row) + [None] * 7)[:7]
        if not identifier or not themes:
            continue
        if str(language or "").startswith("amazighe") or language in excluded_languages:
            continue
        codes = {
            t.strip().casefold() for t in str(themes).replace(",", ";").split(";") if t.strip()
        }
        truth[str(identifier)] = {
            "themes": sorted(codes),
            "tonality": str(tonality).strip().casefold() if tonality else None,
        }
    return truth


PROVISIONAL_BASE = {
    "fr": "sur le jeu de test fictif ({n} contributions annotées par Claude, qui les a aussi rédigées : évaluation provisoire, en partie circulaire)",
    "ar": "على مجموعة الاختبار الافتراضية ({n} مساهمة صنّفها Claude الذي حررها أيضاً: تقييم مؤقت ودائري جزئياً)",
}
REFERENCE_BASE = {
    "fr": "sur {n} contributions fictives classées par le professeur (amazighe exclu) : évaluation de référence",
    "ar": "على {n} مساهمة افتراضية صنّفها الأستاذ (باستثناء الأمازيغية): التقييم المرجعي",
}
ANONYMISATION_BASE = {
    "fr": "sur le jeu de test fictif ({n} pièges)",
    "ar": "على مجموعة الاختبار الافتراضية ({n} فخاً)",
}
