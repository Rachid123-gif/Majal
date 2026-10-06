"""Import of contributions (CSV, Excel) and their processing: anonymisation, analysis by the
local model (or keywords), location by the gazetteer. Processing is idempotent: a contribution
already analysed with the same model is skipped unless `force`."""

import csv
import io
from collections.abc import Callable
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.orm import Session

from app.config_loader.taxonomy import Taxonomy, load_taxonomy
from app.models import Consultation, Contribution, StudyArea
from app.services.citizens.analyze import analyze
from app.services.citizens.anonymize import Masked, apply, cached_config, find_spans
from app.services.citizens.fallback import cached_analysis_config
from app.services.citizens.gazetteer import Gazetteer, load_gazetteer
from app.services.llm import LLMError, get_provider
from app.settings import get_settings

COLUMNS = {
    "identifiant": "external_id",
    "texte": "original_text",
    "langue": "declared_language",
    "commune": "declared_commune",
    "date": "submitted_on",
    "canal": "channel",
}


class ImportError_(ValueError):
    """Readable message for a non-developer (wrong columns, empty file…)."""


def read_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
        header = [str(c or "").strip().casefold() for c in rows[0]] if rows else []
        records = [dict(zip(header, row, strict=False)) for row in rows[1:]]
    elif filename.lower().endswith(".csv"):
        text_ = content.decode("utf-8-sig")
        records = list(csv.DictReader(io.StringIO(text_)))
        records = [{(k or "").strip().casefold(): v for k, v in r.items()} for r in records]
    else:
        raise ImportError_("Format non reconnu : utilisez le modèle CSV ou Excel (.xlsx).")
    if not records or "texte" not in records[0]:
        raise ImportError_(
            "La colonne « texte » est introuvable. Utilisez le modèle "
            "docs/modeles/contributions-modele.xlsx (colonnes : identifiant, texte, langue, "
            "commune, date, canal)."
        )
    out = []
    for number, record in enumerate(records, start=1):
        text_ = str(record.get("texte") or "").strip()
        if not text_:
            continue
        row: dict[str, Any] = {"original_text": text_}
        for column, field_ in COLUMNS.items():
            value = record.get(column)
            if field_ == "original_text" or value in (None, ""):
                continue
            if field_ == "submitted_on":
                value = value if isinstance(value, date) else date.fromisoformat(str(value)[:10])
            row[field_] = value if isinstance(value, date) else str(value).strip()
        row.setdefault("external_id", f"L{number:04d}")
        out.append(row)
    if not out:
        raise ImportError_("Le fichier ne contient aucune contribution (colonne « texte » vide).")
    return out


def import_contributions(
    session: Session,
    study_area: StudyArea,
    code: str,
    title: str,
    badge: str,
    rows: list[dict[str, Any]],
    source_file: str | None,
    user: str | None,
    method_note: str | None = None,
) -> Consultation:
    consultation = session.scalars(
        select(Consultation).where(
            Consultation.study_area_id == study_area.id, Consultation.code == code
        )
    ).one_or_none()
    if consultation is None:
        consultation = Consultation(study_area_id=study_area.id, code=code)
        session.add(consultation)
    consultation.title = title
    consultation.badge = badge
    consultation.source_file = source_file
    consultation.imported_by = user
    consultation.method_note = method_note
    session.flush()
    for row in rows:
        values = {"consultation_id": consultation.id, "badge": badge, **row}
        statement = insert(Contribution).values(values)
        session.execute(
            statement.on_conflict_do_update(
                index_elements=["consultation_id", "external_id"],
                set_={
                    k: statement.excluded[k]
                    for k in values
                    if k not in ("consultation_id", "external_id")
                },
            )
        )
    session.commit()
    return consultation


def taxonomy_for(study_area: StudyArea) -> Taxonomy:
    root = get_settings().config_dir
    return load_taxonomy(root / "taxonomy" / f"{study_area.taxonomy_profile}.yaml")


def _mask_names(text_: str, names: list[str], placeholder: str) -> str:
    for name in sorted({n.strip() for n in names if len(n.strip()) >= 2}, key=len, reverse=True):
        text_ = text_.replace(name, placeholder)
    return text_


def process_contribution(
    contribution: Contribution,
    gazetteer: Gazetteer,
    taxonomy: Taxonomy,
    provider: Any,
) -> None:
    settings = get_settings()
    anon_config = cached_config(settings.config_dir / "citizens" / "anonymisation.yaml")
    analysis_config = cached_analysis_config(settings.config_dir / "citizens" / "analyse.yaml")
    protected = gazetteer.protected_names()
    original = contribution.original_text
    spans = find_spans(original, anon_config, protected)
    masked_text = apply(original, spans, anon_config)

    result = analyze(masked_text, taxonomy, analysis_config, provider)
    # Names the rules missed, reported by the model: masked if found verbatim, outside places.
    placeholder = anon_config.placeholders["person"].fr
    extra: list[Masked] = []
    for name in result.names:
        name = name.strip()
        if len(name) < 2 or any(name.casefold() in p.casefold() for p in protected):
            continue
        start = original.find(name)
        while start != -1:
            end = start + len(name)
            if not any(start < m.end and m.start < end for m in spans + extra):
                extra.append(Masked("person", start, end, "model"))
            start = original.find(name, end)
    if extra:
        spans = sorted(spans + extra, key=lambda m: m.start)
        masked_text = apply(original, spans, anon_config)
    translation = result.translation_fr
    if translation:
        translation = _mask_names(
            translation, [original[m.start : m.end] for m in extra], placeholder
        )

    location = gazetteer.locate(
        masked_text, result.place_fr or result.place, contribution.declared_commune
    )
    counts: dict[str, int] = {}
    for item in spans:
        counts[item.kind] = counts.get(item.kind, 0) + 1
    contribution.anonymized_text = masked_text
    contribution.anonymization = {"counts": counts, "items": [m.to_json() for m in spans]}
    contribution.language = result.language
    contribution.translation_fr = translation
    contribution.themes = result.themes
    contribution.tonality = result.tonality
    contribution.place_text = location.matched or result.place
    contribution.place_id = location.place_id
    contribution.territory_id = location.territory_id
    contribution.analysis = {
        "mode": result.mode,
        "model": result.model,
        "duration_s": result.duration_s,
        "error": result.error,
        "location_method": location.method,
        "model_place": result.place,
        "keywords": result.keywords,
        "translation": "automatic" if result.language != "fr" else "original",
    }


def run(
    session: Session,
    consultation: Consultation,
    force: bool = False,
    progress: Callable[[int, int, Contribution], None] | None = None,
) -> dict[str, int]:
    study_area = session.get(StudyArea, consultation.study_area_id)
    assert study_area is not None
    taxonomy = taxonomy_for(study_area)
    analysis_config = cached_analysis_config(
        get_settings().config_dir / "citizens" / "analyse.yaml"
    )
    gazetteer = load_gazetteer(
        session, study_area.id, analysis_config.place_common_words, analysis_config.place_cues
    )
    try:
        provider = get_provider(get_settings())
    except LLMError:
        provider = None
    model = getattr(provider, "model", None)
    contributions = session.scalars(
        select(Contribution)
        .where(Contribution.consultation_id == consultation.id)
        .order_by(Contribution.external_id)
    ).all()
    stats = {"total": len(contributions), "processed": 0, "skipped": 0, "ai": 0, "keywords": 0}
    for index, contribution in enumerate(contributions, start=1):
        done = contribution.analysis.get("mode") if contribution.analysis else None
        if not force and done == "ai" and contribution.analysis.get("model") == model:
            stats["skipped"] += 1
            continue
        process_contribution(contribution, gazetteer, taxonomy, provider)
        session.commit()
        stats["processed"] += 1
        stats[contribution.analysis["mode"]] += 1
        if progress:
            progress(index, len(contributions), contribution)
    return stats


def fictitious_rows(path: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    import yaml

    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    rows = [
        {
            "external_id": c["id"],
            "original_text": c["text"],
            "declared_commune": c.get("commune_declaree"),
            "channel": "fictif",
        }
        for c in data["contributions"]
    ]
    return data["consultation"], rows
