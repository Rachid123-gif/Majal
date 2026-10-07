"""Data request notes, one per institution, in Word and PDF (stage 5.4).

The text comes from a FIXED template (config/report_templates/note_demande.yaml) completed with
the module's data; nothing is written by the language model. One page of letter, one page of
technical annex. No note may suggest an existing partnership, and citizen contributions are
never cited (both checked against the template's `forbidden` list).
"""

import html
import io
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pydantic import Field

from app.config_loader.data_holders import DataHolders
from app.config_loader.territory import StrictModel, Text
from app.config_loader.validation import load_model

PETROL = RGBColor(0x12, 0x34, 0x3B)
SLATE = RGBColor(0x4A, 0x55, 0x60)


class AnnexColumns(StrictModel):
    priority: Text
    data: Text
    detail: Text
    format: Text
    period: Text
    effect: Text


class NoteTemplate(StrictModel):
    template_version: Text
    status: Text
    language: Text
    header: Text
    header_placeholder: Text
    sender: list[Text]
    place_date: Text
    recipient: list[Text]
    subject: Text
    salutation: Text
    paragraphs: list[Text] = Field(min_length=1)
    closing: Text
    signature: Text
    attachment: Text
    annex_title: Text
    annex_intro: Text
    annex_columns: AnnexColumns
    themes_effect: Text
    context_effect: Text
    footer: Text
    forbidden: list[str] = Field(default_factory=list)


def load_note_template(path: Path) -> NoteTemplate:
    return load_model(path, NoteTemplate, {"paragraphs": '- "Les programmes…"'})


def _plural(n: int, word: str) -> str:
    return word if n == 1 else f"{word}s"


def _de(verb: str) -> str:
    return f"d'{verb}" if verb[:1].lower() in "aeiouyéèêh" else f"de {verb}"


def effects_sentence(
    institution: dict[str, Any], data: dict[str, Any], holders: DataHolders
) -> str:
    """« Ces données permettraient de calculer 3 indicateurs aujourd'hui non disponibles et de
    fiabiliser 1 indicateur aujourd'hui estimé ou issu de données ouvertes. »"""
    verbs = {k: v["verb"]["fr"] for k, v in data["effects"].items()}
    effects = institution["effects"]
    parts: list[str] = []
    n = len(effects["computed"])
    if n:
        text = (
            f"{verbs['computed']} {n} {_plural(n, 'indicateur')} aujourd'hui "
            f"non {_plural(n, 'disponible')}"
        )
        partners = [holders.institution(c) for c in institution["computed_with"]]
        names = [(p.with_data_of or p.name).fr for p in partners if p is not None]
        if names:
            text += " (en complément des données " + " et ".join(names) + ")"
        parts.append(text)
    n = len(effects["reliable"])
    if n:
        parts.append(
            f"{verbs['reliable']} {n} {_plural(n, 'indicateur')} aujourd'hui "
            f"{_plural(n, 'estimé')} ou {_plural(n, 'issu')} de données ouvertes"
        )
    n = len(effects["finer"])
    if n:
        own = [r for r in data["requests"] if r["code"] in institution["requests"]]
        scales = [r["finer_scale"]["fr"] for r in own if r["effects"]["finer"] and r["finer_scale"]]
        scale = f" à l'échelle {scales[0]}" if scales else ""
        parts.append(f"{verbs['finer']} {n} {_plural(n, 'indicateur')}{scale}")
    if institution["themes"]:
        n = len(institution["themes"])
        labels = ", ".join(data["themes"][t]["fr"].lower() for t in institution["themes"])
        parts.append(f"mesurer {n} {_plural(n, 'thème')} aujourd'hui sans indicateur ({labels})")
    own_requests = [r for r in data["requests"] if r["code"] in institution["requests"]]
    if any(r["context"] for r in own_requests):
        parts.append("situer les projets programmés par rapport aux déficits mesurés")
    if not parts:
        return ""
    phrases = [_de(p) for p in parts]
    joined = phrases[0] if len(phrases) == 1 else ", ".join(phrases[:-1]) + " et " + phrases[-1]
    return f"Ces données permettraient {joined}."


NBSP = "\u00a0"


def french_spacing(text: str) -> str:
    """Non-breaking spaces of French typography: never « alone at the end of a line, never a
    colon at the start of one."""
    return (
        text.replace("« ", f"«{NBSP}")
        .replace(" »", f"{NBSP}»")
        .replace(" :", f"{NBSP}:")
        .replace(" ;", f"{NBSP};")
    )


def build_note(
    data: dict[str, Any], institution_code: str, holders: DataHolders, template: NoteTemplate
) -> dict[str, Any]:
    """The note as plain data (letter, then annex rows), shared by the Word and PDF writers."""
    institution = next(i for i in data["institutions"] if i["code"] == institution_code)
    name = institution["name"]["fr"]
    territory = holders.note_territory.fr
    values = {
        "institution": name,
        "territory": territory,
        "effects": effects_sentence(institution, data, holders),
    }

    def fill(text: str) -> str:
        return french_spacing(text.format(**values).replace("  ", " ").strip())

    labels = {i["code"]: i["label"]["fr"] for i in data["indicators"]}
    rows = []
    for request in data["requests"]:  # already ranked: priority first
        if request["code"] not in institution["requests"]:
            continue
        effect_lines = []
        for key in ("computed", "reliable", "finer"):
            codes = request["effects"][key]
            if codes:
                verb = data["effects"][key]["verb"]["fr"].capitalize()
                scale = (
                    f" (à l'échelle {request['finer_scale']['fr']})"
                    if key == "finer" and request["finer_scale"]
                    else ""
                )
                effect_lines.append(f"{verb}{scale} : " + " ; ".join(labels[c] for c in codes))
        if request["themes"]:
            effect_lines.append(
                template.themes_effect.format(
                    themes=", ".join(data["themes"][t]["fr"] for t in request["themes"])
                )
            )
        if request["context"]:
            effect_lines.append(template.context_effect)
        rows.append(
            {
                "priority": request["priority_label"]["fr"],
                "data": french_spacing(request["data"]["fr"]),
                "detail": french_spacing(request["detail"]["fr"]),
                "format": french_spacing(request["format"]["fr"]),
                "period": french_spacing(request["frequency"]["fr"]),
                "effect": [french_spacing(line) for line in effect_lines],
            }
        )
    return {
        "institution": institution_code,
        "header": template.header,
        "header_placeholder": template.header_placeholder,
        "sender": template.sender,
        "place_date": template.place_date,
        "recipient": [fill(line) for line in template.recipient],
        "subject": fill(template.subject),
        "salutation": fill(template.salutation),
        "paragraphs": [fill(p) for p in template.paragraphs],
        "closing": fill(template.closing),
        "signature": template.signature,
        "attachment": fill(template.attachment),
        "annex_title": template.annex_title,
        "annex_intro": fill(template.annex_intro),
        "columns": template.annex_columns.model_dump(),
        "rows": rows,
        "footer": template.footer,
    }


def note_text(note: dict[str, Any]) -> str:
    """Every word of the note, for the checks (forbidden words, placeholders)."""
    parts = [
        note["header"],
        note["header_placeholder"],
        *note["sender"],
        note["place_date"],
        *note["recipient"],
        note["subject"],
        note["salutation"],
        *note["paragraphs"],
        note["closing"],
        note["signature"],
        note["attachment"],
        note["annex_title"],
        note["annex_intro"],
        note["footer"],
    ]
    for row in note["rows"]:
        parts += [row["data"], row["detail"], row["format"], row["period"], *row["effect"]]
    return "\n".join(parts)


# ------------------------------------------------------------------------------------- Word


def _footer(section: Any, text: str) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.runs[0].font.size = Pt(8)
    paragraph.runs[0].font.italic = True
    paragraph.runs[0].font.color.rgb = SLATE


def _para(
    document: Any,
    text: str,
    size: float = 10,
    bold: bool = False,
    italic: bool = False,
    color: RGBColor | None = None,
    after: float = 4,
    align: Any = None,
) -> Any:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.space_before = Pt(0)
    if align is not None:
        paragraph.alignment = align
    return paragraph


def _shade(cell: Any, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def to_docx(note: dict[str, Any]) -> bytes:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2.2)
    _footer(section, note["footer"])
    style = document.styles["Normal"]
    style.font.name = "IBM Plex Sans"
    style.font.size = Pt(10)

    _para(document, note["header"], 11, bold=True, color=PETROL, after=0)
    _para(document, note["header_placeholder"], 8.5, italic=True, color=SLATE, after=10)
    for line in note["sender"]:
        _para(document, line, 9.5, after=0)
    _para(document, note["place_date"], 9.5, after=8, align=WD_ALIGN_PARAGRAPH.RIGHT)
    for line in note["recipient"]:
        _para(document, line, 9.5, after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _para(document, "", 4, after=6)
    _para(document, note["subject"], 10, bold=True, after=8)
    _para(document, note["salutation"], 10, after=6)
    for paragraph in note["paragraphs"]:
        p = _para(document, paragraph, 10, after=5)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _para(document, note["closing"], 10, after=14)
    _para(document, note["signature"], 10, after=8, align=WD_ALIGN_PARAGRAPH.RIGHT)
    attachment = _para(document, note["attachment"], 8.5, italic=True, color=SLATE, after=0)
    attachment.runs[0].add_break(WD_BREAK.PAGE)

    _para(document, note["annex_title"], 13, bold=True, color=PETROL, after=4)
    _para(document, note["annex_intro"], 8.5, color=SLATE, after=6)
    keys = ["priority", "data", "detail", "format", "period", "effect"]
    widths = [Cm(1.9), Cm(4.4), Cm(2.4), Cm(2.4), Cm(2.0), Cm(3.5)]
    table = document.add_table(rows=1, cols=len(keys))
    table.style = "Table Grid"
    for index, key in enumerate(keys):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        cell.text = note["columns"][key]
        run = cell.paragraphs[0].runs[0]
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xF6, 0xF3, 0xEC)
        _shade(cell, "12343B")
    for row in note["rows"]:
        cells = table.add_row().cells
        for index, key in enumerate(keys):
            cells[index].width = widths[index]
            value = "\n".join(row[key]) if key == "effect" else row[key]
            cells[index].text = value
            for paragraph in cells[index].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(7.5)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# -------------------------------------------------------------------------------------- PDF


def to_html(note: dict[str, Any]) -> str:
    e = html.escape
    rows = "".join(
        "<tr>"
        + "".join(
            f"<td>{e(row[k])}</td>" for k in ("priority", "data", "detail", "format", "period")
        )
        + "<td>"
        + "<br>".join(e(line) for line in row["effect"])
        + "</td></tr>"
        for row in note["rows"]
    )
    columns = note["columns"]
    head = "".join(
        f"<th>{e(columns[k])}</th>"
        for k in ("priority", "data", "detail", "format", "period", "effect")
    )
    return f"""<!doctype html><html lang="fr"><head><meta charset="utf-8"><style>
@page {{ size: A4; margin: 18mm 22mm 20mm;
  @bottom-center {{ content: "{e(note["footer"])}";
    font: italic 8pt 'IBM Plex Sans'; color: #4a5560; }} }}
body {{ font-family: 'IBM Plex Sans', 'Noto Sans', sans-serif; font-size: 10pt;
  color: #12343b; line-height: 1.38; }}
p {{ margin: 0 0 5pt; }}
.header {{ font-weight: 600; font-size: 11pt; margin: 0; }}
.placeholder {{ font-style: italic; font-size: 8.5pt; color: #4a5560; margin-bottom: 10pt; }}
.block p {{ margin: 0; font-size: 9.5pt; }}
.right {{ text-align: right; }}
.subject {{ font-weight: 600; margin: 10pt 0 8pt; }}
.body p {{ text-align: justify; }}
.closing {{ margin-top: 6pt; margin-bottom: 14pt; }}
.small {{ font-size: 8.5pt; color: #4a5560; font-style: italic; }}
.annex {{ page-break-before: always; }}
h2 {{ font-family: 'EB Garamond', serif; font-weight: 500; font-size: 15pt; margin: 0 0 4pt; }}
.intro {{ font-size: 8.5pt; color: #4a5560; margin-bottom: 6pt; }}
table {{ border-collapse: collapse; width: 100%; font-size: 7.5pt; line-height: 1.3; }}
th {{ background: #12343b; color: #f6f3ec; text-align: left; padding: 3pt 4pt; }}
td {{ border: 0.5pt solid #c9c3b6; padding: 3pt 4pt; vertical-align: top; }}
</style></head><body>
<p class="header">{e(note["header"])}</p>
<p class="placeholder">{e(note["header_placeholder"])}</p>
<div class="block">{"".join(f"<p>{e(line)}</p>" for line in note["sender"])}</div>
<p class="right" style="margin-top:6pt">{e(note["place_date"])}</p>
<div class="block right">{"".join(f"<p>{e(line)}</p>" for line in note["recipient"])}</div>
<p class="subject">{e(note["subject"])}</p>
<p>{e(note["salutation"])}</p>
<div class="body">{"".join(f"<p>{e(p)}</p>" for p in note["paragraphs"])}</div>
<p class="closing">{e(note["closing"])}</p>
<p class="right">{e(note["signature"])}</p>
<p class="small">{e(note["attachment"])}</p>
<section class="annex">
<h2>{e(note["annex_title"])}</h2>
<p class="intro">{e(note["annex_intro"])}</p>
<table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>
</section></body></html>"""


def to_pdf(note: dict[str, Any]) -> bytes:
    from weasyprint import HTML

    pdf: bytes = HTML(string=to_html(note)).write_pdf()
    return pdf


def pdf_pages(note: dict[str, Any]) -> int:
    from weasyprint import HTML

    return len(HTML(string=to_html(note)).render().pages)
