"""Word (DOCX) and PDF exports of a report (French; Arabic right-to-left exports at stage 7)."""

import html
import io
from datetime import datetime
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

WATERMARK = {
    "fr": "Document de travail généré par MAJAL — à valider par un urbaniste",
    "ar": "وثيقة عمل أنتجها «مجال» — في انتظار مصادقة مختص في التعمير",
}
MODE_NOTE = {
    "ai": "Rédaction : intelligence artificielle locale ({model}), chiffres insérés et contrôlés par MAJAL.",
    "mixed": "Rédaction : intelligence artificielle locale ({model}) ; certaines sections rédigées automatiquement sans IA.",
    "fallback": "Rédaction automatique sans intelligence artificielle (modèle indisponible) à partir des mêmes faits.",
}
STATUS = {"brouillon": "Brouillon", "relu": "Relu", "valide": "Validé"}
PETROL = RGBColor(0x12, 0x34, 0x3B)
TERRACOTTA = RGBColor(0xA8, 0x46, 0x1F)
SLATE = RGBColor(0x4A, 0x55, 0x60)


def _date(value: str | None) -> str:
    if not value:
        return ""
    return datetime.fromisoformat(value).strftime("%d/%m/%Y")


def key_indicator_rows(report: dict[str, Any]) -> list[tuple[str, str, str, str]]:
    """(indicator, value, year, source) for every value fact of the sheet."""
    rows = []
    for fact in report["fact_sheet"].get("facts", []):
        if fact["kind"] == "value":
            rows.append(
                (
                    fact["label"]["fr"],
                    fact["text"]["fr"],
                    str(fact.get("year") or ""),
                    fact.get("source") or "",
                )
            )
    return rows


def _shade(cell: Any, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def to_docx(report: dict[str, Any], map_png: bytes | None) -> bytes:
    content = report["content"]
    document = Document()
    for section in document.sections:
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin = Cm(2.5)
        header = section.header.paragraphs[0]
        header.text = "MAJAL — مجال  ·  " + content.get("grid_label", "")
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = SLATE
        if report["status"] != "valide":
            mark = section.header.add_paragraph(WATERMARK["fr"])
            mark.alignment = WD_ALIGN_PARAGRAPH.CENTER
            mark.runs[0].font.size = Pt(9)
            mark.runs[0].font.bold = True
            mark.runs[0].font.color.rgb = TERRACOTTA
        footer = section.footer.paragraphs[0]
        footer.text = (
            f"MAJAL · Rapport généré le {_date(report.get('finished_at'))} · "
            f"{STATUS.get(report['status'], report['status'])} · Démonstrateur"
        )
        footer.runs[0].font.size = Pt(8)
        footer.runs[0].font.color.rgb = SLATE

    style = document.styles["Normal"]
    style.font.name = "IBM Plex Sans"
    style.font.size = Pt(10.5)

    title = document.add_paragraph()
    run = title.add_run(f"{content['title']} — {content['identity']['name_fr']}")
    run.font.size = Pt(24)
    run.font.color.rgb = PETROL
    run.font.name = "EB Garamond"
    sub = document.add_paragraph(f"{content['identity']['description_fr']}")
    sub.runs[0].font.color.rgb = SLATE
    note = document.add_paragraph(
        MODE_NOTE.get(report.get("writing_mode") or "", "").format(model=report["model"])
    )
    note.runs[0].font.size = Pt(8.5)
    note.runs[0].font.color.rgb = SLATE
    note2 = document.add_paragraph(
        f"{content.get('grid_label', '')} — {content.get('evaluation_label', '')}"
    )
    note2.runs[0].font.size = Pt(8.5)
    note2.runs[0].font.color.rgb = TERRACOTTA

    if map_png:
        document.add_picture(io.BytesIO(map_png), width=Cm(16))

    for section in content["sections"]:
        heading = document.add_heading(f"{section['number']}. {section['title']}", level=2)
        for run in heading.runs:
            run.font.color.rgb = PETROL
        for paragraph in section["paragraphs"]:
            document.add_paragraph(paragraph["text"])
        if section["code"] == "sources":
            for line in section.get("sources", []):
                document.add_paragraph(line, style="List Bullet")
            for line in section.get("notes", []):
                p = document.add_paragraph(line)
                p.runs[0].font.size = Pt(9)
        if section["code"] == "synthese":
            rows = key_indicator_rows(report)
            if rows:
                document.add_heading("Indicateurs clés", level=3)
                table = document.add_table(rows=1, cols=4)
                table.style = "Table Grid"
                for cell, label in zip(
                    table.rows[0].cells, ("Indicateur", "Valeur", "Année", "Source"), strict=True
                ):
                    cell.text = label
                    _shade(cell, "12343B")
                    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xF6, 0xF3, 0xEC)
                    cell.paragraphs[0].runs[0].font.bold = True
                for row in rows:
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row, strict=True):
                        cell.text = value
                        cell.paragraphs[0].runs[0].font.size = Pt(8.5)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def to_html(report: dict[str, Any], map_data_uri: str | None) -> str:
    content = report["content"]
    e = html.escape
    parts = []
    for section in content["sections"]:
        body = "".join(f"<p>{e(p['text'])}</p>" for p in section["paragraphs"])
        if section["code"] == "sources":
            body += (
                "<ul>"
                + "".join(f"<li>{e(line)}</li>" for line in section.get("sources", []))
                + "</ul>"
            )
            body += "".join(f"<p class='small'>{e(line)}</p>" for line in section.get("notes", []))
        if section["code"] == "synthese":
            rows = key_indicator_rows(report)
            if rows:
                body += "<h3>Indicateurs clés</h3><table><tr><th>Indicateur</th><th>Valeur</th><th>Année</th><th>Source</th></tr>"
                body += "".join(
                    f"<tr><td>{e(a)}</td><td class='num'>{e(b)}</td><td>{e(c)}</td><td>{e(d)}</td></tr>"
                    for a, b, c, d in rows
                )
                body += "</table>"
        parts.append(
            f"<section><h2>{section['number']}. {e(section['title'])}</h2>{body}</section>"
        )
    watermark = (
        f"<div class='watermark'>{e(WATERMARK['fr'])}</div>" if report["status"] != "valide" else ""
    )
    image = f"<img class='map' src='{map_data_uri}' alt=''>" if map_data_uri else ""
    mode = e(MODE_NOTE.get(report.get("writing_mode") or "", "").format(model=report["model"]))
    return f"""<!doctype html><html lang="fr"><head><meta charset="utf-8"><style>
@page {{ size: A4; margin: 2.4cm 2cm 2cm 2cm;
  @top-left {{ content: "MAJAL — مجال · {e(content.get("grid_label", ""))}"; font: 8pt 'IBM Plex Sans'; color: #4a5560; }}
  @bottom-left {{ content: "MAJAL · généré le {_date(report.get("finished_at"))} · {STATUS.get(report["status"], "")} · Démonstrateur"; font: 8pt 'IBM Plex Sans'; color: #4a5560; }}
  @bottom-right {{ content: counter(page) " / " counter(pages); font: 8pt 'IBM Plex Sans'; color: #4a5560; }} }}
body {{ font-family: 'IBM Plex Sans', 'Noto Sans', 'Noto Naskh Arabic', sans-serif; font-size: 10.5pt; color: #12343b; line-height: 1.5; }}
h1 {{ font-family: 'EB Garamond', serif; font-weight: 500; font-size: 26pt; margin: 0; }}
h2 {{ font-family: 'EB Garamond', serif; font-weight: 500; font-size: 16pt; margin: 18pt 0 6pt; page-break-after: avoid; }}
h3 {{ font-size: 11pt; margin-top: 12pt; }}
.sub {{ color: #4a5560; margin: 2pt 0 8pt; }}
.note {{ font-size: 8.5pt; color: #4a5560; margin: 2pt 0; }}
.eval {{ font-size: 8.5pt; color: #a8461f; margin: 2pt 0 10pt; }}
.small {{ font-size: 8.5pt; color: #4a5560; }}
.map {{ width: 100%; margin: 8pt 0; }}
table {{ border-collapse: collapse; width: 100%; font-size: 8.5pt; }}
th {{ background: #12343b; color: #f6f3ec; text-align: left; padding: 4pt; }}
td {{ border-bottom: 0.5pt solid #ddd; padding: 3pt 4pt; vertical-align: top; }}
td.num {{ white-space: nowrap; }}
.watermark {{ position: fixed; top: 45%; left: -10%; width: 120%; text-align: center; transform: rotate(-30deg);
  font-size: 22pt; color: rgba(168, 70, 31, 0.13); font-weight: 600; z-index: -1; }}
</style></head><body>{watermark}
<h1>{e(content["title"])} — {e(content["identity"]["name_fr"])}</h1>
<p class="sub">{e(content["identity"]["description_fr"])}</p>
<p class="note">{mode}</p>
<p class="eval">{e(content.get("grid_label", ""))} — {e(content.get("evaluation_label", ""))}</p>
{image}{"".join(parts)}</body></html>"""


def to_pdf(report: dict[str, Any], map_png: bytes | None) -> bytes:
    import base64

    from weasyprint import HTML

    uri = f"data:image/png;base64,{base64.b64encode(map_png).decode()}" if map_png else None
    pdf: bytes = HTML(string=to_html(report, uri)).write_pdf()
    return pdf
