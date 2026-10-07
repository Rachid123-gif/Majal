"""Stage 5.5: summary spreadsheet of the data requests."""

import io

from openpyxl import load_workbook

from app.services.data_needs.excel import to_xlsx
from tests.test_data_notes import DATA


def test_spreadsheet_is_ranked_with_priority_and_the_three_verbs() -> None:
    workbook = load_workbook(io.BytesIO(to_xlsx(DATA, "l'agglomération de Rabat")))
    assert workbook.sheetnames == ["Lisez-moi", "Données demandées", "Institutions"]
    sheet = workbook["Données demandées"]
    headers = [c.value for c in sheet[1]]
    for column in ("Rang", "Priorité", "Calculer", "Fiabiliser", "Affiner", "Période", "Fréquence"):
        assert column in headers, column
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(DATA["requests"])
    priority = headers.index("Priorité")
    assert rows[0][priority] == "Essentielle"
    assert rows[-1][priority] == "Contexte"
    data_column = headers.index("Donnée demandée")
    assert str(rows[0][data_column]).startswith("Limites officielles")


def test_spreadsheet_never_mentions_citizen_contributions() -> None:
    workbook = load_workbook(io.BytesIO(to_xlsx(DATA, "Rabat")))
    text = " ".join(
        str(value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows(values_only=True)
        for value in row
        if value
    ).casefold()
    assert "contribution" not in text and "fictive" not in text
