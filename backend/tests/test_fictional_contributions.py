"""The fictitious citizen contributions (data/fictif/rabat/contributions.yaml) follow the rules of
the project owner: no official function, no accusation, no political, religious or insulting
content (config/citizens/termes-interdits.yaml), no unknown proper name, and they follow the
neutral random plan (data/fictif/rabat/plan.csv) instead of the real indicators."""

import csv
import re
import unicodedata
from collections import Counter
from typing import Any

import pytest
import yaml

from app.settings import REPO_ROOT

DATA = REPO_ROOT / "data" / "fictif" / "rabat"
SET = yaml.safe_load((DATA / "contributions.yaml").read_text(encoding="utf-8"))
CONTRIBUTIONS: list[dict[str, Any]] = SET["contributions"]
PLAN = {row["id"]: row for row in csv.DictReader((DATA / "plan.csv").open(encoding="utf-8"))}
FORBIDDEN = yaml.safe_load(
    (REPO_ROOT / "config" / "citizens" / "termes-interdits.yaml").read_text(encoding="utf-8")
)
TAXONOMY = yaml.safe_load(
    (REPO_ROOT / "config" / "taxonomy" / "urbain.yaml").read_text(encoding="utf-8")
)
AR_LETTER = "ء-ي"


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFC", text).replace("’", "'")
    return re.sub("[ً-ْـ]", "", text).casefold()


def pattern(term: str, arabic: bool) -> re.Pattern[str]:
    t = re.escape(normalize(term))
    if arabic:
        return re.compile(
            rf"(?<![{AR_LETTER}])[وفبلك]?(?:ال)?{t}(?:ه|ها|ة|ت|ا|ات|ان|ين)?(?![{AR_LETTER}])"
        )
    return re.compile(rf"(?<![\w-]){t}(?![\w-])")


def forbidden_hits(text: str) -> list[str]:
    hits = []
    normalized = normalize(text)
    for category, languages in FORBIDDEN.items():
        for language, terms in languages.items():
            for term in terms:
                if pattern(term, language == "ar").search(normalized):
                    hits.append(f"{category}: {term}")
    return hits


def test_the_set_is_complete_and_matches_the_random_plan() -> None:
    assert SET["consultation"]["badge"] == "fictitious"
    ids = [c["id"] for c in CONTRIBUTIONS]
    assert len(ids) == len(set(ids)) == len(PLAN) == 400
    for c in CONTRIBUTIONS:
        plan = PLAN[c["id"]]
        assert c["unit"] == plan["unit"], c["id"]
        assert "|".join(c["themes"]) == plan["themes"], c["id"]
        assert c["language"] == plan["language"], c["id"]
        assert c["tonality"] == plan["tonality"], c["id"]
        assert (c["place"] is not None) == (plan["place_cited"] == "True"), c["id"]
        assert ("commune_declaree" in c) == (plan["commune_declared"] == "True"), c["id"]
        assert bool(c["pii"]) == (plan["pii_trap"] == "True"), c["id"]
        assert c.get("reference_sample", False) == (plan["reference_sample"] == "True"), c["id"]
        if c["language"] == "amazigh_latin":
            assert "approximative" in c["note"], c["id"]


def test_theme_weights_come_from_the_editable_file_and_read_no_indicator() -> None:
    config = yaml.safe_load(
        (REPO_ROOT / "config" / "citizens" / "jeu-fictif.yaml").read_text(encoding="utf-8")
    )
    assert sum(config["theme_weights"].values()) == 100
    assert len(config["reference_sample"]) == 30
    script = (REPO_ROOT / "scripts" / "plan_fictional_contributions.py").read_text()
    code = script.split('"""', 2)[2]  # without the docstring
    for source in ("config/indicators", "diagnostic", "indicator", "/api/", "psycopg"):
        assert source not in code, source


def test_annotations_use_the_taxonomy() -> None:
    themes = {t["code"] for t in TAXONOMY["themes"]}
    for c in CONTRIBUTIONS:
        assert set(c["themes"]) <= themes, c["id"]
        assert c["tonality"] in TAXONOMY["tonalities"], c["id"]


def test_traps_are_in_the_text_and_visibly_fictitious() -> None:
    for c in CONTRIBUTIONS:
        for item in c["pii"]:
            assert item["value"] in c["text"], (c["id"], item)
            digits = re.sub(r"\D", "", item["value"])
            if item["type"] == "phone":  # 06 00 00 0x xx: never a plausible real number
                assert digits.removeprefix("212").removeprefix("0").startswith("600000"), c["id"]
            if item["type"] == "national_id":
                assert item["value"].startswith("ZZ"), c["id"]
            if item["type"] == "email":
                assert item["value"].endswith("@example.com"), c["id"]
    assert sum(len(c["pii"]) for c in CONTRIBUTIONS) >= 80


@pytest.mark.parametrize("contribution", CONTRIBUTIONS, ids=lambda c: c["id"])
def test_no_forbidden_term(contribution: dict[str, Any]) -> None:
    assert forbidden_hits(contribution["text"]) == []


def test_the_forbidden_list_really_catches(  # the test above has a meaning
) -> None:
    assert forbidden_hits("Le président de l'arrondissement n'a rien fait.")
    assert forbidden_hits("القائد لم يتدخل")
    assert forbidden_hits("kayn rachwa f lidara")
    assert forbidden_hits("C'est de la corruption.")
    assert not forbidden_hits("Le trottoir principal de la partie basse est abîmé.")
    assert not forbidden_hits("الطريق الرئيسي مليء بالحفر")


# Latin-script proper names: only the places cited, the fictitious first names of the traps,
# and ordinary words that start a sentence or are capitalised by convention.
ALLOWED_CAPITALISED = {
    "Avenue",
    "Centre",
    "Harhoura",
    "Bravo",
    "Merci",
    "Choukran",
    "Contact",
    "Mon",
    "Ma",
    "Je",
    "Il",
    "Les",
    "Le",
    "La",
    "Un",
    "Une",
    "À",
    "Pour",
    "Sur",
    "Dans",
    "Depuis",
    "Nous",
    "Pourrait-on",
    "Écrivez-moi",
    "Coupures",
    "Smiti",
    "Bghina",
    "Nqtar7o",
    "Tobis",
    "Drari",
    "Triq",
    "Lma",
    "Zan9a",
    "Wlidi",
    "Bnti",
    "L9ism",
    "L7di9a",
    "Tumubilin",
    "Nra",
    "Tarwa",
    "Ur",
    "F",
    "Rue",
    "Route",
    "Place",
    "Hay",
    "Secteur",
    "El",
    "Al",
    "Bab",
    "Hadchi",
    "CIN",
    "Rabat",
    "Salé",
}


def test_no_unknown_proper_name_in_latin_script() -> None:
    known: set[str] = set(ALLOWED_CAPITALISED)
    for c in CONTRIBUTIONS:
        for value in [c.get("place") or "", c.get("unit") or ""] + [
            p["value"]
            for p in c["pii"]  # fictitious names, addresses and numbers of the traps
        ]:
            known.update(re.findall(r"[\wÀ-ÿ'-]+", value))
    unknown: Counter[str] = Counter()
    for c in CONTRIBUTIONS:
        for match in re.finditer(r"(?<![\w'])[A-ZÀ-Ý][\wÀ-ÿ'-]+", c["text"]):
            word, before = match.group(0), c["text"][: match.start()].rstrip()
            if not before or before[-1] in ".!?;:":
                continue  # starts a sentence
            if word not in known and word.split("'")[-1] not in known:
                unknown[f"{c['id']}:{word}"] += 1
    assert not unknown, sorted(unknown)


def test_the_annotated_sheet_still_matches_the_set() -> None:
    """The texts given for annotation (docs/evaluation/annotation-professeur.xlsx) are never
    rewritten: the annotation stays valid when the set grows."""
    from openpyxl import load_workbook

    sheet = load_workbook(
        REPO_ROOT / "docs" / "evaluation" / "annotation-professeur.xlsx", read_only=True
    )["À classer"]
    by_id = {c["id"]: c for c in CONTRIBUTIONS}
    rows = [r for r in sheet.iter_rows(min_row=2, values_only=True) if r[1]]
    assert len(rows) == 30
    for row in rows:
        assert by_id[row[1]]["text"] == row[3], row[1]
        assert by_id[row[1]].get("reference_sample") is True, row[1]
