"""Files derived from data/fictif/rabat/contributions.yaml (stage 4).

Run: `backend/.venv/bin/python scripts/build_citizen_files.py`
- data/fictif/rabat/contributions.csv : the fictitious set in the import format;
- docs/modeles/contributions-modele.xlsx and .csv : empty import template;
- docs/evaluation/annotation-professeur.xlsx : 30 contributions to classify by the professor
  (reference evaluation, independent from Claude's provisional annotation).
"""

import csv
import random
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "fictif" / "rabat" / "contributions.yaml"
TAXONOMY = ROOT / "config" / "taxonomy" / "urbain.yaml"
IMPORT_COLUMNS = ["identifiant", "texte", "langue", "commune", "date", "canal"]
LANGUAGE_LABELS = {
    "fr": "français",
    "ar": "arabe",
    "darija_ar": "darija (arabe)",
    "darija_latin": "darija (latin)",
    "amazigh_latin": "amazighe (latin)",
}
PETROL = PatternFill("solid", fgColor="12343B")
WHITE = Font(color="F6F3EC", bold=True)
SAMPLE_SIZE = 30
SEED = 30


def header(sheet, columns: list[str], widths: list[int]) -> None:  # type: ignore[no-untyped-def]
    sheet.append(columns)
    for cell, width in zip(sheet[1], widths, strict=True):
        cell.fill, cell.font = PETROL, WHITE
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.freeze_panes = "A2"


def main() -> None:
    data = yaml.safe_load(SOURCE.read_text(encoding="utf-8"))
    taxonomy = yaml.safe_load(TAXONOMY.read_text(encoding="utf-8"))
    contributions = data["contributions"]

    # 1. Fictitious set in the import format.
    out = ROOT / "data" / "fictif" / "rabat" / "contributions.csv"
    with out.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(IMPORT_COLUMNS)
        for c in contributions:
            writer.writerow(
                [c["id"], c["text"], "", c.get("commune_declaree", ""), "", "fictif"]
            )

    # 2. Empty import template.
    models = ROOT / "docs" / "modeles"
    models.mkdir(parents=True, exist_ok=True)
    with (models / "contributions-modele.csv").open("w", newline="", encoding="utf-8") as h:
        csv.writer(h).writerow(IMPORT_COLUMNS)
    book = Workbook()
    sheet = book.active
    sheet.title = "Contributions"
    header(sheet, IMPORT_COLUMNS, [14, 80, 14, 22, 14, 16])
    sheet.append(["C-001", "Exemple : le jardin du quartier n'est plus entretenu.", "", "", "", ""])
    notes = book.create_sheet("Mode d'emploi")
    for line in [
        "Une ligne par contribution. Seule la colonne « texte » est obligatoire.",
        "identifiant : libre et unique (sinon MAJAL en attribue un).",
        "langue : facultative (détectée automatiquement).",
        "commune : facultative, nom de l'arrondissement ou de la commune tel que l'a indiqué l'habitant.",
        "date : facultative, AAAA-MM-JJ. canal : facultatif (réunion, formulaire, fictif…).",
        "Ne mettez aucune donnée personnelle inutile : MAJAL masque noms, téléphones, numéros de",
        "carte d'identité, adresses précises, e-mails et plaques avant tout traitement par l'IA.",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 110
    book.save(models / "contributions-modele.xlsx")

    # 3. Sheet for the professor: 30 contributions without traps, stratified by language.
    # Written once: the professor fills it in, it is never overwritten.
    sheet_path = ROOT / "docs" / "evaluation" / "annotation-professeur.xlsx"
    if sheet_path.exists():
        print(f"{len(contributions)} contributions ; fiche du professeur conservée.")
        return
    rng = random.Random(SEED)
    pool = [c for c in contributions if not c["pii"]]
    by_language: dict[str, list[dict[str, object]]] = {}
    for c in pool:
        by_language.setdefault(c["language"], []).append(c)
    quotas = {"fr": 10, "ar": 8, "darija_ar": 5, "darija_latin": 5, "amazigh_latin": 2}
    sample = [c for lang, n in quotas.items() for c in rng.sample(by_language[lang], n)]
    rng.shuffle(sample)
    assert len(sample) == SAMPLE_SIZE

    book = Workbook()
    sheet = book.active
    sheet.title = "À classer"
    header(
        sheet,
        ["N°", "Identifiant", "Langue", "Texte", "Thème(s)", "Tonalité", "Lieu cité"],
        [5, 12, 16, 70, 34, 16, 26],
    )
    for number, c in enumerate(sample, start=1):
        sheet.append([number, c["id"], LANGUAGE_LABELS[c["language"]], c["text"], "", "", ""])
        sheet.cell(row=number + 1, column=4).alignment = Alignment(wrap_text=True, vertical="top")
    tonalities = DataValidation(
        type="list",
        formula1='"' + ",".join(taxonomy["tonalities"]) + '"',
        allow_blank=True,
    )
    sheet.add_data_validation(tonalities)
    tonalities.add(f"F2:F{SAMPLE_SIZE + 1}")

    help_sheet = book.create_sheet("Thèmes et consignes")
    header(help_sheet, ["Code à écrire", "Thème", "Ce qui en relève"], [22, 46, 80])
    for theme in taxonomy["themes"]:
        help_sheet.append([theme["code"], theme["label"]["fr"], theme["description"]])
    help_sheet.append([])
    for line in [
        "Thème(s) : un ou plusieurs codes de la colonne A, séparés par un point-virgule "
        "(ex. : voirie;eclairage).",
        "Tonalité : demande, plainte, proposition ou satisfaction (liste déroulante).",
        "Lieu cité : le quartier, la rue ou la place tel qu'écrit dans le texte ; vide si aucun.",
        "Contributions fictives, rédigées pour la démonstration : elles ne reflètent pas l'opinion "
        "réelle des habitants.",
    ]:
        help_sheet.append(["", line])
    evaluation = ROOT / "docs" / "evaluation"
    evaluation.mkdir(parents=True, exist_ok=True)
    book.save(evaluation / "annotation-professeur.xlsx")
    print(f"{len(contributions)} contributions ; échantillon de {len(sample)} pour le professeur.")


if __name__ == "__main__":
    main()
